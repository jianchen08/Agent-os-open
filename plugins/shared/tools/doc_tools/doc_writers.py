"""文档写出器——Block[] → md/docx/pdf。

成熟模式移植（不自己造生成器）：
- Word：python-docx 官方 API（add_heading/add_paragraph/add_table）
- PDF：reportlab 官方 API（SimpleDocTemplate + Paragraph/Table flowables；
  中文经内置 CID 字体 STSong-Light，无需外部字体文件）

[来源: python-docx 官方文档（Document.add_heading/add_paragraph/add_table）；
 reportlab 官方文档（platypus SimpleDocTemplate / UnicodeCIDFont）]
"""

from __future__ import annotations

from pathlib import Path

from doc_core import Block, DocumentError, parse_inline

# ── Markdown 写出 ─────────────────────────────────────────────


def blocks_to_markdown(blocks: list[Block]) -> str:
    """Block[] → Markdown 文本（标题/段落/列表/代码/引用/表格）。"""
    lines: list[str] = []
    for b in blocks:
        if b.kind == "heading":
            lines.append("#" * max(1, min(b.level, 6)) + " " + b.text)
        elif b.kind == "paragraph":
            lines.append(b.text)
        elif b.kind == "bullet":
            lines.append("  " * max(0, b.level - 1) + "- " + b.text)
        elif b.kind == "ordered":
            lines.append("  " * max(0, b.level - 1) + "1. " + b.text)
        elif b.kind == "code":
            lines.append(f"```{b.info}\n{b.text}\n```")
        elif b.kind == "quote":
            lines.append("> " + b.text)
        elif b.kind == "table":
            lines.extend(_table_to_markdown(b.rows))
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def _table_to_markdown(rows: list[list[str]]) -> list[str]:
    """表格行 → Markdown 表格（首行为表头）。"""
    if not rows:
        return []
    width = max(len(r) for r in rows)
    out: list[str] = []
    header = list(rows[0]) + [""] * (width - len(rows[0]))
    out.append("| " + " | ".join(header) + " |")
    out.append("| " + " | ".join("---" for _ in range(width)) + " |")
    for row in rows[1:]:
        padded = list(row) + [""] * (width - len(row))
        out.append("| " + " | ".join(padded) + " |")
    return out


# ── Word 写出 ─────────────────────────────────────────────────


def write_docx(blocks: list[Block], path: Path) -> None:
    """Block[] → .docx（python-docx；粗体/斜体/代码内联样式保留）。"""
    from docx import Document

    try:
        doc = Document()
        for b in blocks:
            _add_docx_block(doc, b)
        doc.save(str(path))
    except OSError as e:
        raise DocumentError(f"写入 Word 文件失败：{path.name}（{e}）") from e


def _add_docx_block(doc: object, b: Block) -> None:
    """单个块 → python-docx 段落/表格。"""
    if b.kind == "heading":
        para = doc.add_paragraph(style=f"Heading {max(1, min(b.level, 9))}")  # type: ignore[attr-defined]
        _add_docx_rich_text(para, b.text)
    elif b.kind == "paragraph":
        _add_docx_rich_text(doc.add_paragraph(), b.text)  # type: ignore[attr-defined]
    elif b.kind == "bullet":
        _add_docx_rich_text(doc.add_paragraph(style="List Bullet"), b.text)  # type: ignore[attr-defined]
    elif b.kind == "ordered":
        _add_docx_rich_text(doc.add_paragraph(style="List Number"), b.text)  # type: ignore[attr-defined]
    elif b.kind == "code":
        para = doc.add_paragraph()  # type: ignore[attr-defined]
        run = para.add_run(b.text)
        run.font.name = "Courier New"
    elif b.kind == "quote":
        para = doc.add_paragraph()  # type: ignore[attr-defined]
        run = para.add_run(b.text)
        run.italic = True
    elif b.kind == "table" and b.rows:
        width = max(len(r) for r in b.rows)
        table = doc.add_table(rows=len(b.rows), cols=width)  # type: ignore[attr-defined]
        for i, row in enumerate(b.rows):
            for j in range(width):
                table.cell(i, j).text = row[j] if j < len(row) else ""


def _add_docx_rich_text(para: object, text: str) -> None:
    """带内联样式的文本 → 段落 runs（粗体/斜体/粗斜体/代码）。"""
    for seg_text, style in parse_inline(text):
        run = para.add_run(seg_text)  # type: ignore[attr-defined]
        if style in ("bold", "bolditalic"):
            run.bold = True
        if style in ("italic", "bolditalic"):
            run.italic = True
        if style == "code":
            run.font.name = "Courier New"


# ── PDF 写出 ──────────────────────────────────────────────────

_PDF_FONT = "STSong-Light"  # reportlab 内置 CID 字体（中文支持，无需外部字体文件）
_PDF_HEADING_SIZES = {1: 18, 2: 15, 3: 13, 4: 12, 5: 11, 6: 11}


def write_pdf(blocks: list[Block], path: Path) -> None:
    """Block[] → .pdf（reportlab；中文经内置 CID 字体）。"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    if _PDF_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(_PDF_FONT))

    body = ParagraphStyle("body", fontName=_PDF_FONT, fontSize=10.5, leading=16, spaceAfter=6)
    quote = ParagraphStyle(
        "quote", parent=body, leftIndent=12, textColor=colors.HexColor("#555555")
    )
    code_style = ParagraphStyle(
        "code", parent=body, fontName="Courier", fontSize=9, leading=13, leftIndent=12
    )
    heading_styles = {
        n: ParagraphStyle(
            f"h{n}",
            fontName=_PDF_FONT,
            fontSize=_PDF_HEADING_SIZES[n],
            leading=_PDF_HEADING_SIZES[n] + 6,
            spaceBefore=12,
            spaceAfter=6,
        )
        for n in range(1, 7)
    }

    flowables: list[object] = []
    for b in blocks:
        if b.kind == "heading":
            flowables.append(
                Paragraph(_to_pdf_markup(b.text), heading_styles[max(1, min(b.level, 6))])
            )
        elif b.kind == "paragraph":
            flowables.append(Paragraph(_to_pdf_markup(b.text), body))
        elif b.kind in ("bullet", "ordered"):
            flowables.append(
                ListFlowable(
                    [ListItem(Paragraph(_to_pdf_markup(b.text), body))],
                    bulletType="bullet" if b.kind == "bullet" else "1",
                    leftIndent=18,
                )
            )
        elif b.kind == "code":
            flowables.append(Paragraph(_to_pdf_markup(b.text).replace("\n", "<br/>"), code_style))
            flowables.append(Spacer(1, 6))
        elif b.kind == "quote":
            flowables.append(Paragraph(_to_pdf_markup(b.text), quote))
        elif b.kind == "table" and b.rows:
            flowables.append(_pdf_table(b.rows, body, Table, TableStyle, colors))
            flowables.append(Spacer(1, 6))

    try:
        doc = SimpleDocTemplate(
            str(path),
            pagesize=A4,
            leftMargin=2 * cm,
            rightMargin=2 * cm,
            topMargin=2 * cm,
            bottomMargin=2 * cm,
            title=path.stem,
        )
        doc.build(flowables)  # type: ignore[arg-type]
    except OSError as e:
        raise DocumentError(f"写入 PDF 文件失败：{path.name}（{e}）") from e


def _pdf_table(rows: list[list[str]], body: object, table_cls: object, style_cls: object, colors: object) -> object:
    """表格行 → reportlab Table（含网格线）。"""
    from xml.sax.saxutils import escape

    from reportlab.platypus import Paragraph

    width = max(len(r) for r in rows)
    data = [
        [Paragraph(escape(str(cell)), body) for cell in (list(row) + [""] * (width - len(row)))]  # type: ignore[arg-type]
        for row in rows
    ]
    table = table_cls(data)  # type: ignore[operator]
    table.setStyle(
        style_cls(  # type: ignore[operator]
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),  # type: ignore[attr-defined]
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),  # type: ignore[attr-defined]
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return table


def _to_pdf_markup(text: str) -> str:
    """Markdown 内联标记 → reportlab 段落标记（XML 转义 + b/i 标签）。"""
    from xml.sax.saxutils import escape

    parts: list[str] = []
    for seg_text, style in parse_inline(text):
        esc = escape(seg_text)
        if style == "bold":
            parts.append(f"<b>{esc}</b>")
        elif style == "italic":
            parts.append(f"<i>{esc}</i>")
        elif style == "bolditalic":
            parts.append(f"<b><i>{esc}</i></b>")
        else:
            parts.append(esc)
    return "".join(parts)
